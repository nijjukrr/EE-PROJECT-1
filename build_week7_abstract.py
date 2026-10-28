import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_week7_abstract():
    doc = docx.Document()
    
    # Page setup: Letter size (8.5 x 11.0 inches = 612 x 792 pt) matching original PDF
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    
    # Margins: In original PDF, text x0=90pt (1.25 in), x1=541.5pt (612-541.5 = 70.5pt ~ 0.98 in)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(1.15)
    section.right_margin = Inches(1.0)
    
    # Add page border: single line, 0.5 pt (sz="4"), 24 pt from page edges
    sectPr = section._sectPr
    pgBorders = parse_xml(
        '<w:pgBorders %s w:offsetFrom="page">'
        '  <w:top w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
        '  <w:left w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
        '  <w:right w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
        '</w:pgBorders>' % nsdecls('w')
    )
    sectPr.append(pgBorders)
    
    def add_run(p, text, bold=False, italic=False, size=Pt(10), font_name="Times New Roman"):
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = size
        run.font.color.rgb = RGBColor(0, 0, 0)
        return run

    # 1. Header Image (College Banner)
    header_img_path = os.path.abspath("extracted_header.jpeg")
    if os.path.exists(header_img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.line_spacing = 1.0
        run_img = p_img.add_run()
        run_img.add_picture(header_img_path, width=Inches(6.25))

    # 2. DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(3)
    p_dept.paragraph_format.space_after = Pt(5)
    p_dept.paragraph_format.line_spacing = 1.0
    add_run(p_dept, "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", bold=True, size=Pt(11))

    # 3. Two-column metadata table (CLASS/SEM, DATE, BATCH, PROJECT GUIDE)
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    for row in meta_table.rows:
        row.cells[0].width = Inches(3.2)
        row.cells[1].width = Inches(3.1)
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                '<w:tcMar %s>'
                '  <w:top w:w="0" w:type="dxa"/>'
                '  <w:bottom w:w="0" w:type="dxa"/>'
                '  <w:left w:w="0" w:type="dxa"/>'
                '  <w:right w:w="0" w:type="dxa"/>'
                '</w:tcMar>' % nsdecls('w')
            )
            tcPr.append(tcMar)

    # Row 0: CLASS /SEM and DATE
    p00 = meta_table.rows[0].cells[0].paragraphs[0]
    p00.paragraph_format.space_before = Pt(1)
    p00.paragraph_format.space_after = Pt(2)
    p00.paragraph_format.line_spacing = 1.0
    add_run(p00, "CLASS /SEM : ", bold=True, size=Pt(10.5))
    add_run(p00, "II / III", bold=False, size=Pt(10.5))

    p01 = meta_table.rows[0].cells[1].paragraphs[0]
    p01.paragraph_format.space_before = Pt(1)
    p01.paragraph_format.space_after = Pt(2)
    p01.paragraph_format.line_spacing = 1.0
    add_run(p01, "DATE : ", bold=True, size=Pt(10.5))
    add_run(p01, "10.10.2026", bold=False, size=Pt(10.5))

    # Row 1: BATCH and PROJECT GUIDE
    p10 = meta_table.rows[1].cells[0].paragraphs[0]
    p10.paragraph_format.space_before = Pt(1)
    p10.paragraph_format.space_after = Pt(4)
    p10.paragraph_format.line_spacing = 1.0
    add_run(p10, "BATCH        : ", bold=True, size=Pt(10.5))
    add_run(p10, "03", bold=False, size=Pt(10.5))

    p11 = meta_table.rows[1].cells[1].paragraphs[0]
    p11.paragraph_format.space_before = Pt(1)
    p11.paragraph_format.space_after = Pt(4)
    p11.paragraph_format.line_spacing = 1.0
    add_run(p11, "PROJECT GUIDE: ", bold=True, size=Pt(10.5))
    add_run(p11, "Mrs. GAYATHRI V AP/CSE", bold=False, size=Pt(10.5))

    # 4. TEAM MEMBERS heading
    p_tm = doc.add_paragraph()
    p_tm.paragraph_format.space_before = Pt(3)
    p_tm.paragraph_format.space_after = Pt(2)
    p_tm.paragraph_format.line_spacing = 1.0
    add_run(p_tm, "TEAM MEMBERS :", bold=True, size=Pt(10.5))

    # Team members in 1-column vertical list (matching original layout exactly)
    members = [
        "1. PADMANABHAN SURESH BABU (714025104188)",
        "2. MITHUNEESH K (714025104144)",
        "3. NISHANTH KR (714025104173)",
        "4. NITISH RG (714025104183)"
    ]
    for idx, mem in enumerate(members):
        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_before = Pt(0.5)
        p_m.paragraph_format.space_after = Pt(1.5 if idx == len(members)-1 else 0.5)
        p_m.paragraph_format.line_spacing = 1.0
        # In original, member 1 was bold, others regular. Here regular is clean and consistent
        add_run(p_m, mem, bold=False, size=Pt(10))

    # 5. ABSTRACT:
    p_abs_h = doc.add_paragraph()
    p_abs_h.paragraph_format.space_before = Pt(4)
    p_abs_h.paragraph_format.space_after = Pt(1.5)
    p_abs_h.paragraph_format.line_spacing = 1.0
    add_run(p_abs_h, "ABSTRACT:", bold=True, size=Pt(10.5))

    # Project Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0.5)
    p_title.paragraph_format.space_after = Pt(2.5)
    p_title.paragraph_format.line_spacing = 1.0
    add_run(p_title, "BROZER: A PRIVACY BOUNDARY FOR BROWSER AGENTS WITH ON-DEVICE VISUAL PERCEPTION:", bold=True, size=Pt(10))

    # Abstract Paragraph
    p_abs_body = doc.add_paragraph()
    p_abs_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs_body.paragraph_format.space_before = Pt(0)
    p_abs_body.paragraph_format.space_after = Pt(3.5)
    p_abs_body.paragraph_format.line_spacing = 1.10
    add_run(p_abs_body,
        "Autonomous browser agents represent a transformative advance in artificial intelligence, allowing models to interpret "
        "complex web interfaces and autonomously execute multi-step user tasks. However, contemporary browser agents introduce "
        "severe privacy vulnerabilities by capturing unfiltered page viewports and complete DOM hierarchies, which are transmitted "
        "directly to remote vision-language models. This unredacted context often contains sensitive credentials, financial records, "
        "and personal communications, creating significant compliance and surveillance risks. This project presents Brozer, a novel "
        "privacy-first browser agent architecture featuring a strictly enforced local privacy boundary and on-device visual perception. "
        "Built upon the WebBrain runtime, Brozer re-engineers browser agent execution under the principle of client-side containment: "
        "perceive locally, protect locally, and reason externally with minimal sanitized context. A six-stage pipeline (Perceive, Detect, "
        "Sanitize, Reason, Policy Check, and Act) guarantees zero raw credential leakage across the trust boundary.",
        bold=False, size=Pt(10)
    )

    # 6. Progress on 10.10.2026 (Week 7)
    p_p1_h = doc.add_paragraph()
    p_p1_h.paragraph_format.space_before = Pt(3)
    p_p1_h.paragraph_format.space_after = Pt(1.5)
    p_p1_h.paragraph_format.line_spacing = 1.0
    add_run(p_p1_h, "Progress on 10.10.2026 (Week 7)", bold=True, size=Pt(10.5))

    p_p1_b = doc.add_paragraph()
    p_p1_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_p1_b.paragraph_format.space_before = Pt(0)
    p_p1_b.paragraph_format.space_after = Pt(3.5)
    p_p1_b.paragraph_format.line_spacing = 1.10
    add_run(p_p1_b,
        "During this phase, the on-device perception and multi-signal privacy sanitization pipelines were consolidated into the core "
        "browser agent extension. Lightweight DOM semantic extraction and client-side ONNX Runtime Web perception models were implemented "
        "to localize interactive elements and detect sensitive PII across passwords, contact numbers, and payment fields. The context "
        "sanitization engine replaces confidential values with semantic tokens ([EMAIL], [PASSWORD]) and applies visual blurring to "
        "private screen regions prior to network transmission. Additionally, the client-side execution firewall (Policy Gate) was "
        "engineered to enforce deterministic security rules (ALLOW, REDACT, BLOCK, ASK USER) and autofill credentials directly from "
        "the client vault during live testing across real-world web environments.",
        bold=False, size=Pt(10)
    )

    # 7. Progress for 24.10.2026
    p_p2_h = doc.add_paragraph()
    p_p2_h.paragraph_format.space_before = Pt(3)
    p_p2_h.paragraph_format.space_after = Pt(1.5)
    p_p2_h.paragraph_format.line_spacing = 1.0
    add_run(p_p2_h, "Progress for 24.10.2026", bold=True, size=Pt(10.5))

    p_p2_b = doc.add_paragraph()
    p_p2_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_p2_b.paragraph_format.space_before = Pt(0)
    p_p2_b.paragraph_format.space_after = Pt(18)
    p_p2_b.paragraph_format.line_spacing = 1.10
    add_run(p_p2_b,
        "During the upcoming phase, comprehensive end-to-end benchmarking, privacy verification, and institutional documentation "
        "will be completed. The system's zero-leakage guarantee will be formally evaluated against benchmark web task suites (e-commerce, "
        "social platforms, and authentication workflows) to measure latency overhead, perception accuracy, and redaction precision. "
        "Furthermore, final academic deliverables will be prepared, including the comprehensive project report, research manuscript, "
        "and technical presentation slides. Software packaging, environment dependency freezing, and user operation guides will be "
        "finalized to ensure seamless replication and demonstration during final project evaluation.",
        bold=False, size=Pt(10)
    )

    # 8. Signature line table (PROJECT GUIDE, PROJECT COORDINATOR, HOD)
    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    
    col_widths = [Inches(2.1), Inches(2.3), Inches(1.9)]
    alignments = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]
    sig_texts = ["PROJECT GUIDE", "PROJECT COORDINATOR", "HOD"]
    
    for idx, cell in enumerate(sig_table.rows[0].cells):
        cell.width = col_widths[idx]
        p_sig = cell.paragraphs[0]
        p_sig.alignment = alignments[idx]
        p_sig.paragraph_format.space_before = Pt(14)
        p_sig.paragraph_format.space_after = Pt(0)
        p_sig.paragraph_format.line_spacing = 1.0
        add_run(p_sig, sig_texts[idx], bold=True, size=Pt(10.5))
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            '<w:tcMar %s>'
            '  <w:top w:w="0" w:type="dxa"/>'
            '  <w:bottom w:w="0" w:type="dxa"/>'
            '  <w:left w:w="0" w:type="dxa"/>'
            '  <w:right w:w="0" w:type="dxa"/>'
            '</w:tcMar>' % nsdecls('w')
        )
        tcPr.append(tcMar)

    out_docx = "WEEK 7 NEW ABSTRACT- final.docx"
    doc.save(out_docx)
    print(f"Saved {out_docx}")

if __name__ == "__main__":
    create_week7_abstract()
