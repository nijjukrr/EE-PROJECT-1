import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_week8_abstract():
    doc = docx.Document()
    
    # Page setup: Standard A4 (8.27 x 11.69 inches) matching Indian academic standards
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    
    # Outer margins:
    section.top_margin = Inches(0.46)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    
    # Add page border: single line, 0.5 pt (sz="4"), 24 pt offset from page edge
    sectPr = section._sectPr
    pgBorders = parse_xml(
        '<w:pgBorders %s w:offsetFrom="page">'
        '  <w:top w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
        '  <w:left w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
        '  <w:right w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
        '</w:pgBorders>' % nsdecls('w')
    )
    sectPr.append(pgBorders)
    
    def add_run(p, text, bold=False, italic=False, size=Pt(10), font_name="Times New Roman"):
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = size
        run.font.color.rgb = RGBColor(0, 0, 0)
        return run

    # 1. Header Banner (Exact original soft-toned 3D header, pure white background)
    header_img = os.path.abspath("header_trimmed.png")
    if os.path.exists(header_img):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(0)
        p_img.paragraph_format.line_spacing = 1.0
        r_img = p_img.add_run()
        r_img.add_picture(header_img, width=Inches(6.35))

    # 2. DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(44)
    p_dept.paragraph_format.space_after = Pt(20)
    p_dept.paragraph_format.line_spacing = 1.0
    add_run(p_dept, "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", bold=True, size=Pt(11.5))

    # 3. Two-Column Metadata Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    col_w = Inches(3.17)
    for row in meta_table.rows:
        row.cells[0].width = col_w
        row.cells[1].width = col_w
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                '<w:tcMar %s>'
                '  <w:top w:w="0" w:type="dxa"/>'
                '  <w:bottom w:w="0" w:type="dxa"/>'
                '  <w:left w:w="0" w:type="dxa"/>'
                '  <w:right w:w="0" w:type="dxa"/>'
                '</w:tcMar>' % nsdecls('w')
            )
            tcPr.append(tcMar)

    # Row 0: CLASS/SEM & DATE (Week 8: 24.10.2026)
    p00 = meta_table.rows[0].cells[0].paragraphs[0]
    p00.paragraph_format.space_before = Pt(0)
    p00.paragraph_format.space_after = Pt(2)
    p00.paragraph_format.line_spacing = 1.0
    add_run(p00, "CLASS/SEM: ", bold=False, size=Pt(10))
    add_run(p00, "II / III", bold=False, size=Pt(10))

    p01 = meta_table.rows[0].cells[1].paragraphs[0]
    p01.paragraph_format.space_before = Pt(0)
    p01.paragraph_format.space_after = Pt(2)
    p01.paragraph_format.line_spacing = 1.0
    add_run(p01, "DATE: ", bold=False, size=Pt(10))
    add_run(p01, "24.10.2026", bold=False, size=Pt(10))

    # Row 1: BATCH & PROJECT GUIDE
    p10 = meta_table.rows[1].cells[0].paragraphs[0]
    p10.paragraph_format.space_before = Pt(2)
    p10.paragraph_format.space_after = Pt(0)
    p10.paragraph_format.line_spacing = 1.0
    add_run(p10, "BATCH: ", bold=False, size=Pt(10))
    add_run(p10, "03", bold=False, size=Pt(10))

    p11 = meta_table.rows[1].cells[1].paragraphs[0]
    p11.paragraph_format.space_before = Pt(2)
    p11.paragraph_format.space_after = Pt(0)
    p11.paragraph_format.line_spacing = 1.0
    add_run(p11, "PROJECT GUIDE: ", bold=False, size=Pt(10))
    add_run(p11, "Gayathri V AP/CSE", bold=False, size=Pt(10))

    # 4. Project Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(22)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.paragraph_format.line_spacing = 1.0
    add_run(p_title, "Brozer— See Nothing It Shouldn't.", bold=True, size=Pt(11.5))

    # 5. Members: Heading
    p_mem_h = doc.add_paragraph()
    p_mem_h.paragraph_format.space_before = Pt(0)
    p_mem_h.paragraph_format.space_after = Pt(14)
    p_mem_h.paragraph_format.line_spacing = 1.0
    add_run(p_mem_h, "Members:", bold=True, size=Pt(10.5))

    # 4 Members Indented with individual line spacing
    members_data = [
        ("1.  ", "Padmanabhan Suresh Babu ", True, "(714025104188)", False),
        ("2.  ", "Mithuneesh K ", False, "(714025104144)", False),
        ("3.  ", "Nitish RG ", False, "(714025204183)", False),
        ("4.  ", "Nishanth KR ", False, "(714025104713)", False)
    ]
    for idx, (num, name, b_name, reg, b_reg) in enumerate(members_data):
        pm = doc.add_paragraph()
        pm.paragraph_format.left_indent = Inches(1.1)
        pm.paragraph_format.space_before = Pt(1)
        pm.paragraph_format.space_after = Pt(20 if idx == len(members_data)-1 else 7.5)
        pm.paragraph_format.line_spacing = 1.15
        add_run(pm, num, bold=False, size=Pt(10))
        add_run(pm, name, bold=b_name, size=Pt(10))
        add_run(pm, reg, bold=b_reg, size=Pt(10))

    # 6. Abstract: Heading
    p_abs_h = doc.add_paragraph()
    p_abs_h.paragraph_format.space_before = Pt(0)
    p_abs_h.paragraph_format.space_after = Pt(8)
    p_abs_h.paragraph_format.line_spacing = 1.0
    add_run(p_abs_h, "Abstract:", bold=True, size=Pt(10.5))

    # Abstract Paragraph with first-line indent
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.paragraph_format.first_line_indent = Inches(0.45)
    p_abs.paragraph_format.space_before = Pt(0)
    p_abs.paragraph_format.space_after = Pt(42)
    p_abs.paragraph_format.line_spacing = 1.18
    add_run(p_abs,
        "During this review period, we completed the integration of the Brozer privacy boundary layer "
        "with the browser extension, reasoning backend, and local execution companion. The tiered detection "
        "pipeline was evaluated to verify detection accuracy for structured sensitive entities and contextual "
        "information across dynamic webpages. We validated the DOM-to-pixel geometry mapping to locate sensitive "
        "elements reliably across responsive browser layouts and fractional device pixel ratios. The redaction "
        "engine successfully enforces local masking to ensure that protected regions never reach the reasoning "
        "server. In addition, the privacy egress gate and log sanitizer were comprehensively tested to confirm "
        "that outgoing communication payloads and system logs remain free of unredacted sensitive values.",
        bold=False, size=Pt(10)
    )

    # 7. Progress for [31.10.2026] Heading (Week 8 upcoming milestones)
    p_prog_h = doc.add_paragraph()
    p_prog_h.paragraph_format.space_before = Pt(0)
    p_prog_h.paragraph_format.space_after = Pt(18)
    p_prog_h.paragraph_format.line_spacing = 1.0
    add_run(p_prog_h, "Progress for [31.10.2026]", bold=True, size=Pt(10.5))

    # 4 Bullet Points for 31.10.2026
    bullets = [
        "Benchmark end-to-end task execution latency across multi-step browser workflows.",
        "Validate policy check enforcement and local autofill action-gating mechanisms.",
        "Measure redaction precision and recall across diverse web forms and authenticated dashboards.",
        "Finalize system performance documentation and prepare comprehensive review deliverables."
    ]
    for b in bullets:
        pb = doc.add_paragraph()
        pb.paragraph_format.left_indent = Inches(0.25)
        pb.paragraph_format.space_before = Pt(2)
        pb.paragraph_format.space_after = Pt(8)
        pb.paragraph_format.line_spacing = 1.15
        add_run(pb, "• ", bold=True, size=Pt(10))
        add_run(pb, b, bold=False, size=Pt(10))

    # 8. Signature Row: Clean digital recreation without pen signatures, dedicated signing space matching master template
    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    
    col_w_sig = [Inches(2.1), Inches(2.3), Inches(1.9)]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]
    sig_names = ["PROJECT GUIDE", "PROJECT COORDINATOR", "HOD"]
    
    for idx, cell in enumerate(sig_table.rows[0].cells):
        cell.width = col_w_sig[idx]
        p_sig = cell.paragraphs[0]
        p_sig.alignment = aligns[idx]
        p_sig.paragraph_format.space_before = Pt(100) # Perfectly calibrated signing gap matching master layout
        p_sig.paragraph_format.space_after = Pt(0)
        p_sig.paragraph_format.line_spacing = 1.0
        add_run(p_sig, sig_names[idx], bold=True, size=Pt(10.5))
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            '<w:tcMar %s>'
            '  <w:top w:w="0" w:type="dxa"/>'
            '  <w:bottom w:w="0" w:type="dxa"/>'
            '  <w:left w:w="0" w:type="dxa"/>'
            '  <w:right w:w="0" w:type="dxa"/>'
            '</w:tcMar>' % nsdecls('w')
        )
        tcPr.append(tcMar)

    out_docx = "WEEK 8 NEW ABSTRACT- final.docx"
    doc.save(out_docx)
    print(f"Saved {out_docx}")

if __name__ == "__main__":
    create_week8_abstract()
