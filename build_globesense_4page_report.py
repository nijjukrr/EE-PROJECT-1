import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_globesense_4page_doc():
    doc = docx.Document()

    # Page setup: Standard A4 (8.27 x 11.69 inches)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

    # Margins matching dbms final report-nitz.pdf
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.50)

    # Page border: single line, 0.75 pt (sz="6"), 24 pt offset from page edge
    sectPr = section._sectPr
    pgBorders = parse_xml(
        '<w:pgBorders %s w:offsetFrom="page">'
        '  <w:top w:val="single" w:sz="6" w:space="24" w:color="000000"/>'
        '  <w:left w:val="single" w:sz="6" w:space="24" w:color="000000"/>'
        '  <w:bottom w:val="single" w:sz="6" w:space="24" w:color="000000"/>'
        '  <w:right w:val="single" w:sz="6" w:space="24" w:color="000000"/>'
        '</w:pgBorders>' % nsdecls('w')
    )
    sectPr.append(pgBorders)

    def add_run(p, text, bold=False, italic=False, size=Pt(11), font_name="Times New Roman"):
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = size
        run.font.color.rgb = RGBColor(0, 0, 0)
        return run

    def add_para(text="", bold=False, size=Pt(11), align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=Pt(0), space_after=Pt(3), line_spacing=1.15):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = space_before
        p.paragraph_format.space_after = space_after
        p.paragraph_format.line_spacing = line_spacing
        if text:
            add_run(p, text, bold=bold, size=size)
        return p

    def add_heading(text, space_before=Pt(6), space_after=Pt(3)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = space_before
        p.paragraph_format.space_after = space_after
        p.paragraph_format.line_spacing = 1.15
        add_run(p, text, bold=True, size=Pt(11))
        return p

    def add_bullet(text, space_after=Pt(2), space_before=Pt(0), line_spacing=1.15):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.20)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_before = space_before
        p.paragraph_format.space_after = space_after
        p.paragraph_format.line_spacing = line_spacing
        add_run(p, "\u2022 ", bold=False, size=Pt(11))
        add_run(p, text, bold=False, size=Pt(11))
        return p

    def add_subbullet(text, space_after=Pt(1.5), space_before=Pt(0), line_spacing=1.15):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.42)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_before = space_before
        p.paragraph_format.space_after = space_after
        p.paragraph_format.line_spacing = line_spacing
        add_run(p, "\u2013 ", bold=False, size=Pt(11))
        add_run(p, text, bold=False, size=Pt(11))
        return p

    # =========================================================================
    # PAGE 1
    # =========================================================================

    # 1. Header Banner
    header_img = os.path.abspath("college_header.jpeg")
    if os.path.exists(header_img):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.line_spacing = 1.0
        r_img = p_img.add_run()
        r_img.add_picture(header_img, width=Inches(6.4))

    # 2. DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(4)
    p_dept.paragraph_format.space_after = Pt(6)
    p_dept.paragraph_format.line_spacing = 1.0
    add_run(p_dept, "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", bold=True, size=Pt(11.5))

    # 3. Two-Column Metadata Table (3 rows)
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_w = Inches(3.55)
    for row in meta_table.rows:
        row.cells[0].width = col_w
        row.cells[1].width = col_w
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                '<w:tcMar %s>'
                '  <w:top w:w="0" w:type="dxa"/>'
                '  <w:bottom w:w="0" w:type="dxa"/>'
                '  <w:left w:w="0" w:type="dxa"/>'
                '  <w:right w:w="0" w:type="dxa"/>'
                '</w:tcMar>' % nsdecls('w')
            )
            tcPr.append(tcMar)

    # Row 0: CLASS/SEM & DATE
    p00 = meta_table.rows[0].cells[0].paragraphs[0]
    p00.paragraph_format.space_before = Pt(0)
    p00.paragraph_format.space_after = Pt(2)
    p00.paragraph_format.line_spacing = 1.15
    add_run(p00, "CLASS/SEM: ", bold=False, size=Pt(11))
    add_run(p00, "B.E.CSE / II YEAR", bold=True, size=Pt(11))

    p01 = meta_table.rows[0].cells[1].paragraphs[0]
    p01.paragraph_format.space_before = Pt(0)
    p01.paragraph_format.space_after = Pt(2)
    p01.paragraph_format.line_spacing = 1.15
    add_run(p01, "DATE: ", bold=False, size=Pt(11))

    # Row 1: PROJECT TYPE & STUDENT
    p10 = meta_table.rows[1].cells[0].paragraphs[0]
    p10.paragraph_format.space_before = Pt(2)
    p10.paragraph_format.space_after = Pt(2)
    p10.paragraph_format.line_spacing = 1.15
    add_run(p10, "PROJECT TYPE: ", bold=False, size=Pt(11))
    add_run(p10, "DBMS PROJECT", bold=True, size=Pt(11))

    p11 = meta_table.rows[1].cells[1].paragraphs[0]
    p11.paragraph_format.space_before = Pt(2)
    p11.paragraph_format.space_after = Pt(2)
    p11.paragraph_format.line_spacing = 1.15
    add_run(p11, "STUDENT: ", bold=False, size=Pt(11))
    add_run(p11, "NISHANTH KR", bold=True, size=Pt(11))

    # Row 2: PROJECT GUIDE & EMPTY
    p20 = meta_table.rows[2].cells[0].paragraphs[0]
    p20.paragraph_format.space_before = Pt(2)
    p20.paragraph_format.space_after = Pt(6)
    p20.paragraph_format.line_spacing = 1.15
    add_run(p20, "PROJECT GUIDE: ", bold=False, size=Pt(11))
    add_run(p20, "Ms.Abishree.S AP/CSE", bold=True, size=Pt(11))

    # 4. Project Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.paragraph_format.line_spacing = 1.15
    add_run(p_title, "PROJECT TITLE: ", bold=False, size=Pt(11))
    add_run(p_title, "GLOBESENSE: REAL-TIME 3D GEOSPATIAL TRACKING AND VISUALIZATION PLATFORM", bold=True, size=Pt(11))

    # 5. Abstract
    add_heading("Abstract", space_before=Pt(4), space_after=Pt(3))

    add_para(
        "GlobeSense is a high-performance, browser-based geospatial intelligence and database-driven visualization "
        "platform that unifies real-time public telemetry streams onto a single interactive 3D virtual globe. Traditional "
        "tracking systems operate within isolated, single-purpose web silos, forcing analysts, students, and operators "
        "to toggle between disparate applications to monitor flights, maritime vessels, orbital satellites, seismic events, "
        "and atmospheric conditions. GlobeSense replaces these fragmented interfaces with an integrated operational "
        "command console, enabling fluid transitions from a planetary macro-overview down to individual aircraft vectors, "
        "naval hulls, and municipal street infrastructure.",
        space_after=Pt(4)
    )

    add_para(
        "The client architecture features multi-layered geospatial data filtering, an interactive click-to-track engine "
        "that computes camera look-at vectors and renders fading kinematic trajectories, a first-person cockpit simulation "
        "over authentic terrain elevation models, and custom post-processing visual sensor modes (Night Vision, Thermal FLIR, "
        "and CRT Phosphor). Analytical capabilities include geodetic distance measurement, polygon surface area calculation, "
        "turn-by-turn multi-modal routing, and meteorological overlays including dynamic wind particle velocity vectors, "
        "precipitation radar, cloud reflectance, lightning density, and tropical cyclone trajectories.",
        space_after=Pt(4)
    )

    add_para(
        "Engineered as a distributed client-server ecosystem, GlobeSense integrates an embedded SQLite relational database "
        "for mission parameters, DuckDB Spatial SQL executing analytical queries over columnar Apache Parquet archives, "
        "and a secure Node.js proxy server that protects upstream API secrets, sanitizes queries, and prevents client-side "
        "performance degradation. The platform ingests telemetry from open public repositories—including OpenSky Network, "
        "adsb.lol, CelesTrak, USGS, NASA FIRMS, NOAA, and OpenStreetMap—operating entirely without mandatory paid API keys. "
        "Optional free keys activate photorealistic 3D urban tilesets and real-time maritime AIS tracking, demonstrating "
        "an accessible and production-grade approach to database-driven geospatial analytics.",
        space_after=Pt(4)
    )

    # 6. Objectives (Part 1 - 5 items on Page 1)
    add_heading("Objectives", space_before=Pt(4), space_after=Pt(3))

    add_bullet("To architect a centralized, database-backed web application unifying heterogeneous live public telemetry feeds onto an interactive 3D virtual globe.")
    add_bullet("To implement a persistent relational database schema for indexing satellites, flights, ground stations, CCTV nodes, and user annotation layers.")
    add_bullet("To integrate spatial SQL query processing via DuckDB and Parquet storage for high-speed spatial predicate evaluation and geospatial boundary filtering.")
    add_bullet("To engineer real-time entity tracking with dynamic camera orientation, fading historical trails, and synchronized live telemetry HUD cards.")
    add_bullet("To build a realistic first-person cockpit simulation mode synchronizing real-time flight vectors with 3D digital elevation models.", space_after=Pt(0))

    # Explicit page break to Page 2
    doc.add_page_break()

    # =========================================================================
    # PAGE 2
    # =========================================================================

    # Objectives (Part 2 - remaining 5 items on Page 2)
    add_bullet("To develop custom GPU-accelerated GLSL shader pipelines providing multi-spectral visual sensor modes (thermal infrared, night vision, and CRT).")
    add_bullet("To integrate real-time atmospheric and meteorological observation layers, including Doppler radar, wind velocity fields, and cyclone forecasting.")
    add_bullet("To implement SGP4 orbital propagation algorithms for real-time satellite trajectory calculation based on CelesTrak Two-Line Element (TLE) sets.")
    add_bullet("To construct a resilient Node.js middleware proxy implementing request rate-limiting, spatial bounding-box clipping, and credential isolation.")
    add_bullet("To deploy the platform on cloud infrastructure with zero mandatory paid subscription requirements, ensuring open educational and research accessibility.", space_after=Pt(6))

    # Work Completed
    add_heading("Work Completed", space_before=Pt(4), space_after=Pt(3))

    add_para(
        "This project is built upon the open-source God's Eye View architecture (MIT License), significantly extended "
        "and customized with robust spatial database pipelines, telemetry ingestion layers, and custom client interfaces. "
        "The core work completed includes:",
        space_after=Pt(4)
    )

    # Bullet 1
    add_bullet("Database Schema & Relational Telemetry Storage:")
    add_subbullet("Designed and deployed the SQLite relational database schema (satellite-operations.sqlite) defining structured tables for satellites, flights, cameras, crew, and data_layers.")
    add_subbullet("Implemented primary key constraints, foreign key referential integrity, and timestamped audit logs for mission telemetry and tracking sessions.")
    add_subbullet("Configured DuckDB Spatial queries over partitioned Apache Parquet datasets to execute spatial filtering (ST_PointOnSurface, ST_Area_Spheroid) and bounding box clipping.")
    add_subbullet("Integrated Overpass QL relational querying over OpenStreetMap datasets to extract regional infrastructure, runways, and military coordinates.")

    # Bullet 2
    add_bullet("Multi-Stream Live Telemetry Ingestion Pipeline:")
    add_subbullet("Implemented automated ingestion handlers for ADS-B flight data (OpenSky Network & adsb.lol REST APIs) with state deduplication.")
    add_subbullet("Integrated CelesTrak TLE orbital ephemeris feeds utilizing SGP4 orbital propagation algorithms to calculate real-time satellite positions.")
    add_subbullet("Integrated USGS real-time seismic event feeds and NASA FIRMS active thermal wildfire hotspot detection data.")
    add_subbullet("Connected live maritime AIS vessel tracking, public CCTV camera feeds, transit networks, and micro-mobility bikeshare data.")

    # Bullet 3
    add_bullet("3D Geospatial Globe & Coordinate System Engine:")
    add_subbullet("Implemented high-performance 3D rendering using CesiumJS, integrating WGS84 ellipsoidal geometry and EGM96 geoid gravitational height corrections.")
    add_subbullet("Configured multi-source basemap switching across Esri World Imagery, OpenStreetMap vector tiles, and Google Photorealistic 3D Tiles via Cesium ion.")
    add_subbullet("Developed custom GLSL post-processing shaders delivering immersive sensor vision modes: Night Vision (green phosphor), Thermal FLIR (false-color heat map), and CRT Tactical Scanlines.")

    # Bullet 4
    add_bullet("Interactive Tracking, Cockpit Simulation & Telemetry Cards:")
    add_subbullet("Built click-to-track camera locking with orientation calculation, heading interpolation, and fading polyline flight trails.")
    add_subbullet("Developed first-person cockpit simulation mode projecting real elevation terrain beneath moving aircraft models.")
    add_subbullet("Created real-time telemetry HUD overlay displaying callsign, aircraft type, altitude, ground speed, heading, and geographic coordinates.", space_after=Pt(0))

    # Explicit page break to Page 3
    doc.add_page_break()

    # =========================================================================
    # PAGE 3
    # =========================================================================

    # Work Completed (Continuation)
    add_bullet("Drawing Tools, Spatial Routing & Mission Recording:", space_after=Pt(1))
    add_subbullet("Implemented geodetic measurement tools for real-time calculation of surface distances, bearings, and enclosed polygon areas.", space_after=Pt(1))
    add_subbullet("Integrated OpenStreetMap Nominatim and OSRM routing engine supporting turn-by-turn driving, walking, and cycling route navigation.", space_after=Pt(1))
    add_subbullet("Developed Scene Director timeline playback engine enabling users to record, script, and replay cinematic geospatial flythrough tours.", space_after=Pt(1))

    add_bullet("Security, Proxy Middleware & Deployment Architecture:", space_after=Pt(1))
    add_subbullet("Engineered a resilient Node.js proxy server enforcing client rate-limiting, error sanitization, and upstream API credential encapsulation.", space_after=Pt(1))
    add_subbullet("Configured automated CORS policy handling and response caching to minimize upstream API quotas and latency.", space_after=Pt(1))
    add_subbullet("Deployed production client and proxy server on Vercel with high-availability edge routing.", space_after=Pt(4))

    # Current Status
    add_heading("Current Status", space_before=Pt(3), space_after=Pt(2))

    add_para(
        "The GlobeSense platform is fully functional and live in production on Vercel. It seamlessly renders 19 interactive "
        "geospatial data layers across the globe, executing real-time entity tracking, live telemetry HUD updates, cockpit "
        "perspective simulation, GLSL multi-spectral visual filters, weather intelligence displays, spatial vector annotation "
        "tools, and turn-by-turn route calculations. The application operates without requiring any paid API subscriptions, "
        "with free optional API keys unlocking photorealistic Google 3D Tiles and maritime vessel tracking.",
        space_after=Pt(3)
    )

    add_para(
        "On the database tier, the embedded SQLite database engine (satellite-operations.sqlite) manages relational records "
        "for satellites, active flights, traffic cameras, flight crews, and layer configurations with full referential integrity. "
        "DuckDB Spatial executes analytical SQL queries against Apache Parquet vector archives, while the Node.js middleware proxy "
        "manages upstream connection pooling, response payload caching, and secure API credential isolation. The platform maintains "
        "sub-second query execution and a fluid 60 FPS rendering rate on modern web browsers.",
        space_after=Pt(3)
    )

    # Technologies Used
    add_heading("Technologies Used", space_before=Pt(3), space_after=Pt(2))

    techs = [
        ("Frontend 3D Engine: ", "CesiumJS (v1.124), WebGL, GLSL Custom Shaders"),
        ("Web Technologies: ", "JavaScript (ES6+), HTML5, CSS3, Vite 6.0 Bundler"),
        ("Database Engines: ", "SQLite 3 (satellite-operations.sqlite), DuckDB Spatial (Parquet SQL)"),
        ("Backend & Middleware: ", "Node.js (Proxy Server), RESTful JSON Endpoints, WebSockets"),
        ("Geospatial & Math Libraries: ", "Satellite.js (SGP4), EGM96-Universal, MGRS, Mapbox Vector Tile"),
        ("External Data Feeds: ", "OpenSky Network, adsb.lol, CelesTrak, USGS, NASA FIRMS, NOAA, OpenStreetMap"),
        ("Map Providers & 3D: ", "Esri World Imagery, OpenStreetMap, Google Photorealistic 3D Tiles (Cesium ion)"),
        ("Hosting & Deployment: ", "Vercel (Production Edge Network Hosting)")
    ]
    for label, val in techs:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(0)
        p_t.paragraph_format.space_after = Pt(1)
        p_t.paragraph_format.line_spacing = 1.15
        add_run(p_t, label, bold=True, size=Pt(11))
        add_run(p_t, val, bold=False, size=Pt(11))

    # Technology Justification
    add_heading("Technology Justification:", space_before=Pt(3), space_after=Pt(2))

    add_para(
        "CesiumJS delivers an authentic WGS84 ellipsoidal globe with true geodetic coordinates, dynamic elevation clamping, "
        "and atmospheric physics that 2D mapping libraries cannot represent. The combination of SQLite and DuckDB Spatial enables "
        "zero-configuration, high-performance relational and columnar spatial SQL processing directly on the server without incurring "
        "the substantial financial and operational overhead of dedicated cloud database instances. Node.js provides non-blocking "
        "asynchronous I/O ideal for proxying high-frequency telemetry streams, while Vite ensures rapid hot-module replacement and "
        "optimized client-side bundle delivery.",
        space_after=Pt(3)
    )

    # Project Plan for Future (Part 1 - 4 items on Page 3)
    add_heading("Project Plan for Future", space_before=Pt(3), space_after=Pt(2))

    add_bullet("Implement historical telemetry replay using time-partitioned SQLite tables to visualize historical flight paths and naval routes.", space_after=Pt(1.5))
    add_bullet("Integrate PostGIS / PostgreSQL cloud database replication for enterprise-grade collaborative multi-user mission tracking.", space_after=Pt(1.5))
    add_bullet("Add real-time spatial geofencing alerts triggered via automated spatial queries when tracked entities enter designated airspace.", space_after=Pt(1.5))
    add_bullet("Incorporate ADS-B Mode-S military aircraft transponder filters and automated anomalous squawk code (7500, 7600, 7700) detection.", space_after=Pt(0))

    # Explicit page break to Page 4
    doc.add_page_break()

    # =========================================================================
    # PAGE 4
    # =========================================================================

    # Project Plan for Future (Part 2 - remaining 4 items on Page 4)
    add_bullet("Expand meteorological modeling by integrating high-resolution GRIB2 weather forecast datasets via WebAssembly decoders.")
    add_bullet("Integrate hands-free speech recognition and AI voice command control for camera navigation and layer toggling.")
    add_bullet("Develop automated PDF and CSV reporting modules for generating situational intelligence briefings and mission logs.")
    add_bullet("Optimize mobile browser touch gestures and responsive HUD layout for mobile and tablet field operation.", space_after=Pt(5))

    # Expected Outcome
    add_heading("Expected Outcome", space_before=Pt(4), space_after=Pt(3))

    add_para(
        "The expected outcome of the GlobeSense project is a production-grade, highly dependable 3D geospatial intelligence platform "
        "that bridges disparate public data streams into a unified situational awareness console. By replacing fragmented tracking "
        "tools with an integrated virtual globe, users gain immediate access to cross-domain activity encompassing aviation, maritime, "
        "orbital, seismic, and meteorological environments.",
        space_after=Pt(4)
    )

    add_para(
        "From a DBMS perspective, the project demonstrates how modern lightweight database architectures—combining SQLite relational "
        "integrity with DuckDB spatial analytics—can process high-velocity geospatial datasets efficiently without prohibitive infrastructure "
        "costs. Administrators and researchers benefit from sub-second coordinate resolution, spatial boundary queries, and secure "
        "credential handling via the proxy layer.",
        space_after=Pt(4)
    )

    add_para(
        "The platform delivers significant educational, analytical, and practical value, empowering students and researchers to study "
        "orbital mechanics, air traffic flow, maritime logistics, and disaster monitoring within an intuitive, open-access browser "
        "environment. By operating entirely on free and open tiers, GlobeSense establishes a sustainable and replicable blueprint for "
        "modern geospatial systems.",
        space_after=Pt(6)
    )

    # Conclusion
    add_heading("Conclusion", space_before=Pt(4), space_after=Pt(3))

    add_para(
        "GlobeSense demonstrates the successful fusion of relational database management, spatial SQL analytics, and cutting-edge 3D WebGL "
        "graphics into an interactive planetary digital twin. By bringing together flights, ships, satellites, earthquakes, wildfires, and "
        "weather phenomena onto a single globe, the system proves that open public data feeds can deliver rich, institutional-grade "
        "intelligence when paired with an optimized backend.",
        space_after=Pt(4)
    )

    add_para(
        "The utilization of CesiumJS, SQLite, DuckDB Spatial, and a secure Node.js proxy establishes a modular, secure, and maintainable "
        "software architecture. GlobeSense exemplifies how contemporary database principles and modern web development can automate "
        "complex geospatial tracking workflows, providing an exceptional and immersive user experience.",
        space_after=Pt(4)
    )

    # Signature Row: 3 columns matching dbms final report-nitz.pdf exactly
    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    col_w_sig = [Inches(2.4), Inches(2.5), Inches(2.3)]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]
    sig_names = ["STUDENT", "PROJECT GUIDE", "HOD"]

    for idx, cell in enumerate(sig_table.rows[0].cells):
        cell.width = col_w_sig[idx]
        p_sig = cell.paragraphs[0]
        p_sig.alignment = aligns[idx]
        p_sig.paragraph_format.space_before = Pt(75) # Dedicated signature signing space
        p_sig.paragraph_format.space_after = Pt(0)
        p_sig.paragraph_format.line_spacing = 1.0
        add_run(p_sig, sig_names[idx], bold=True, size=Pt(11))
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            '<w:tcMar %s>'
            '  <w:top w:w="0" w:type="dxa"/>'
            '  <w:bottom w:w="0" w:type="dxa"/>'
            '  <w:left w:w="0" w:type="dxa"/>'
            '  <w:right w:w="0" w:type="dxa"/>'
            '</w:tcMar>' % nsdecls('w')
        )
        tcPr.append(tcMar)

    out_docx = "globesense_report.docx"
    doc.save(out_docx)
    print(f"Saved {out_docx}")

if __name__ == "__main__":
    create_globesense_4page_doc()
